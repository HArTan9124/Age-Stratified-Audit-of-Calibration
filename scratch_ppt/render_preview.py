"""Approximate PPTX -> PNG previewer (matplotlib). Good enough to catch layout bugs."""
import sys
import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
import matplotlib.patches as mp
from matplotlib.offsetbox import OffsetImage, AnnotationBbox
import matplotlib.image as mpimg
from pptx import Presentation
from pptx.util import Emu
from pptx.enum.shapes import MSO_SHAPE_TYPE

EMU_IN = 914400
SW, SH = 13.333, 7.5

def rgb(c):
    try:
        return "#%02x%02x%02x" % (c[0], c[1], c[2])
    except Exception:
        return None

def draw(prs, path_out):
    n = len(prs.slides)
    fig, axes = plt.subplots(n, 1, figsize=(SW, SH * n))
    if n == 1:
        axes = [axes]
    for ax, slide in zip(axes, prs.slides):
        ax.set_xlim(0, SW); ax.set_ylim(0, SH); ax.invert_yaxis()
        ax.set_aspect("equal"); ax.axis("off")
        ax.add_patch(mp.Rectangle((0, 0), SW, SH, fc="white", ec="none"))
        for sh in slide.shapes:
            if sh.left is None:
                continue
            x, y = sh.left / EMU_IN, sh.top / EMU_IN
            w, h = (sh.width or 0) / EMU_IN, (sh.height or 0) / EMU_IN
            st = sh.shape_type
            if st == MSO_SHAPE_TYPE.PICTURE:
                try:
                    im = mpimg.imread(sh.image.blob and __import__("io").BytesIO(sh.image.blob))
                    ax.imshow(im, extent=(x, x + w, y + h, y), aspect="auto", zorder=5)
                    ax.add_patch(mp.Rectangle((x, y), w, h, fc="none", ec="#bbb", lw=0.5, zorder=6))
                except Exception as e:
                    ax.add_patch(mp.Rectangle((x, y), w, h, fc="#eee", ec="#999", lw=0.5))
                    ax.text(x + w/2, y + h/2, f"[img]\n{e}", ha="center", va="center", fontsize=5)
                continue
            if sh.has_table:
                tbl = sh.table
                nr, nc = len(tbl.rows), len(tbl.columns)
                cw = [c.width / EMU_IN for c in tbl.columns]
                tot = sum(cw) or w
                cw = [c * w / tot for c in cw]
                rh = h / nr
                cy = y
                for ri in range(nr):
                    cx = x
                    for ci in range(nc):
                        cell = tbl.cell(ri, ci)
                        fc = "#1a1a2e" if ri == 0 else ("#ffffff" if ri % 2 else "#f0efea")
                        try:
                            f = cell.fill.fore_color.rgb
                            fc = rgb(f) or fc
                        except Exception:
                            pass
                        ax.add_patch(mp.Rectangle((cx, cy), cw[ci], rh, fc=fc, ec="#ccc", lw=0.4, zorder=7))
                        txt = cell.text
                        _h = fc.lstrip("#")
                        _lum = (int(_h[0:2],16)*0.299 + int(_h[2:4],16)*0.587 + int(_h[4:6],16)*0.114) if len(_h)==6 else 255
                        col = "white" if _lum < 110 else "#1a1a2e"
                        ax.text(cx + 0.05, cy + rh/2, txt, ha="left", va="center",
                                fontsize=7, color=col, zorder=8)
                        cx += cw[ci]
                    cy += rh
                continue
            # autoshape / textbox
            fc = "none"; ec = "none"
            if st == MSO_SHAPE_TYPE.AUTO_SHAPE:
                try:
                    if sh.fill.type is not None:
                        fc = rgb(sh.fill.fore_color.rgb) or "none"
                except Exception:
                    fc = "none"
                try:
                    ec = rgb(sh.line.color.rgb) or "none"
                except Exception:
                    ec = "none"
                ax.add_patch(mp.FancyBboxPatch((x + 0.03, y + 0.03), max(w - 0.06, 0.01), max(h - 0.06, 0.01),
                             boxstyle="round,pad=0.02", fc=fc, ec=ec if ec != "none" else "none",
                             lw=0.8, zorder=2))
            if sh.has_text_frame:
                cy = y + 0.06
                for p in sh.text_frame.paragraphs:
                    if not p.text:
                        cy += 0.16; continue
                    runs = p.runs or []
                    size = 12
                    color = "#1a1a2e"
                    bold = False
                    if runs:
                        r0 = runs[0]
                        if r0.font.size: size = r0.font.size.pt
                        if r0.font.bold: bold = True
                        try:
                            if r0.font.color and r0.font.color.rgb:
                                color = rgb(r0.font.color.rgb)
                        except Exception:
                            pass
                    align = {1: "center", 2: "right"}.get(int(p.alignment) if p.alignment else 1, "left") \
                        if p.alignment is not None else "left"
                    tx = x + (w/2 if align == "center" else (w - 0.1 if align == "right" else 0.1))
                    # wrap
                    import textwrap
                    chars = max(int(w / (size * 0.0095)), 8)
                    lines = []
                    for seg in p.text.split("\n"):
                        lines += textwrap.wrap(seg, chars) or [""]
                    for ln in lines:
                        ax.text(tx, cy, ln, ha=align, va="top", fontsize=size * 0.83,
                                color=color, weight="bold" if bold else "normal", zorder=10)
                        cy += size / 72 * 1.28
                    cy += (p.space_after.pt if p.space_after else 3) / 72
        ax.add_patch(mp.Rectangle((0, 0), SW, SH, fc="none", ec="#333", lw=1))
    plt.subplots_adjust(left=0, right=1, top=1, bottom=0, hspace=0.04)
    fig.savefig(path_out, dpi=110)
    print("wrote", path_out)

if __name__ == "__main__":
    prs = Presentation(sys.argv[1])
    if len(sys.argv) > 3:
        lo, hi = (int(x) for x in sys.argv[3].split("-"))
        keep = [sl for i, sl in enumerate(prs.slides, 1) if lo <= i <= hi]
        from pptx.oxml.ns import qn
        sldIdLst = prs.slides._sldIdLst
        for sid in list(sldIdLst):
            sldIdLst.remove(sid)
        # rebuild a shim object exposing .slides iterable
        class Shim:
            pass
        sh = Shim(); sh.slides = keep
        draw(sh, sys.argv[2]); sys.exit(0)
    draw(prs, sys.argv[2])
