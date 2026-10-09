"""Shared helpers for building the NB05 / NB06 result decks (python-pptx)."""
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from PIL import Image

# ---- palette ---------------------------------------------------------------
INK      = RGBColor(0x1A, 0x1A, 0x2E)   # near-black navy
ACCENT   = RGBColor(0xC0, 0x39, 0x2B)   # ECG red
ACCENT2  = RGBColor(0x2E, 0x5E, 0x8C)   # steel blue
PAPER    = RGBColor(0xFA, 0xFA, 0xF7)
PANEL    = RGBColor(0xF0, 0xEF, 0xEA)
MUTE     = RGBColor(0x5A, 0x5A, 0x66)
WHITE    = RGBColor(0xFF, 0xFF, 0xFF)
LINE     = RGBColor(0xD8, 0xD6, 0xCE)

EMU_W, EMU_H = Inches(13.333), Inches(7.5)


def new_deck():
    prs = Presentation()
    prs.slide_width = EMU_W
    prs.slide_height = EMU_H
    return prs


def _blank(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    bg = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, EMU_W, EMU_H)
    bg.fill.solid(); bg.fill.fore_color.rgb = PAPER
    bg.line.fill.background()
    bg.shadow.inherit = False
    return s


def _tb(slide, x, y, w, h, anchor=MSO_ANCHOR.TOP):
    tb = slide.shapes.add_textbox(x, y, w, h)
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    return tb, tf


def _set(p, text, size, color=INK, bold=False, italic=False, font="Calibri",
         align=PP_ALIGN.LEFT, space_after=4, space_before=0):
    p.text = text
    p.alignment = align
    p.space_after = Pt(space_after)
    p.space_before = Pt(space_before)
    for r in p.runs:
        r.font.size = Pt(size); r.font.bold = bold; r.font.italic = italic
        r.font.name = font; r.font.color.rgb = color


def add_run(p, text, size, color=INK, bold=False, italic=False, font="Calibri"):
    r = p.add_run(); r.text = text
    r.font.size = Pt(size); r.font.bold = bold; r.font.italic = italic
    r.font.name = font; r.font.color.rgb = color
    return r


# ---- title slide ---------------------------------------------------------------
def title_slide(prs, kicker, title, subtitle, meta_lines):
    s = _blank(prs)
    _, tf = _tb(s, Inches(0.9), Inches(0.7), Inches(11.5), Inches(0.5))
    _set(tf.paragraphs[0], kicker, 15, ACCENT, bold=True)

    _, tf = _tb(s, Inches(0.9), Inches(1.15), Inches(11.7), Inches(1.7))
    _set(tf.paragraphs[0], title, 34, INK, bold=True)

    band = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.9), Inches(2.9), Inches(3.2), Inches(0.055))
    band.fill.solid(); band.fill.fore_color.rgb = ACCENT; band.line.fill.background()
    band.shadow.inherit = False

    _, tf = _tb(s, Inches(0.9), Inches(3.1), Inches(11.7), Inches(0.9))
    _set(tf.paragraphs[0], subtitle, 17, MUTE)

    _, tf = _tb(s, Inches(0.9), Inches(4.15), Inches(11.7), Inches(3.1))
    for i, ln in enumerate(meta_lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        _set(p, "", 13)
        add_run(p, "▪  ", 13, ACCENT, bold=True)
        add_run(p, ln, 13, INK)
        p.space_after = Pt(6)
    return s


# ---- content slide -------------------------------------------------------------
def content_slide(prs, index, total, heading, tagline):
    s = _blank(prs)
    # header
    bar = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, EMU_W, Inches(1.02))
    bar.fill.solid(); bar.fill.fore_color.rgb = INK; bar.line.fill.background()
    bar.shadow.inherit = False
    _, tf = _tb(s, Inches(0.55), Inches(0.06), Inches(11.4), Inches(0.92),
                anchor=MSO_ANCHOR.MIDDLE)
    _set(tf.paragraphs[0], heading, 20, WHITE, bold=True)
    if tagline:
        _set(tf.add_paragraph(), tagline, 11, RGBColor(0xC9, 0xC9, 0xD6), space_before=2)
    # page number
    _, tf = _tb(s, Inches(12.2), Inches(0.08), Inches(0.9), Inches(0.9),
                anchor=MSO_ANCHOR.MIDDLE)
    _set(tf.paragraphs[0], f"{index}/{total}", 12, RGBColor(0x9A, 0x9A, 0xA8),
         bold=True, align=PP_ALIGN.RIGHT)
    return s


def bullets(slide, x, y, w, h, items, size=13, gap=6):
    """items: list of (text, level, color, bold)."""
    _, tf = _tb(slide, x, y, w, h)
    for i, it in enumerate(items):
        text, lvl, col, bold = (it + (0, INK, False))[:4] if len(it) < 4 else it
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        _set(p, "", size)
        marker = "▪  " if lvl == 0 else "–  "
        mcol = ACCENT if lvl == 0 else MUTE
        add_run(p, marker, size, mcol, bold=True)
        add_run(p, text, size, col, bold=bold)
        p.space_after = Pt(gap)
        p.space_before = Pt(2 if lvl == 0 and i else 0)
    return tf


def panel(slide, x, y, w, h, fill=PANEL):
    box = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, h)
    box.adjustments[0] = 0.04
    box.fill.solid(); box.fill.fore_color.rgb = fill
    box.line.color.rgb = LINE; box.line.width = Pt(0.75)
    box.shadow.inherit = False
    return box


def panel_title(slide, x, y, w, text, color=ACCENT2):
    _, tf = _tb(slide, x, y, w, Inches(0.4))
    _set(tf.paragraphs[0], text.upper(), 11, color, bold=True)
    return tf


def image_fit(slide, path, x, y, w, h, align="center", valign="middle", border=True):
    """Place image scaled to fit inside the (w,h) box, preserving aspect."""
    iw, ih = Image.open(path).size
    ar = iw / ih
    box_ar = w / h
    if ar > box_ar:
        dw = w; dh = Emu(int(w / ar))
    else:
        dh = h; dw = Emu(int(h * ar))
    dx = x + (0 if align == "left" else (w - dw) if align == "right" else (w - dw) // 2)
    dy = y + (0 if valign == "top" else (h - dh) if valign == "bottom" else (h - dh) // 2)
    pic = slide.shapes.add_picture(path, dx, dy, dw, dh)
    if border:
        pic.line.color.rgb = LINE; pic.line.width = Pt(0.75)
    return pic


def caption(slide, x, y, w, text):
    _, tf = _tb(slide, x, y, w, Inches(0.4))
    _set(tf.paragraphs[0], text, 9.5, MUTE, italic=True)
    return tf


# ---- lightweight table -------------------------------------------------------
def simple_table(slide, x, y, w, rows, col_ratio, header=True, size=10.5,
                 row_h=Inches(0.3)):
    nrow = len(rows); ncol = len(rows[0])
    gt = slide.shapes.add_table(nrow, ncol, x, y, w, row_h * nrow).table
    for ci, r in enumerate(col_ratio):
        gt.columns[ci].width = Emu(int(w * r))
    for ri, row in enumerate(rows):
        for ci, val in enumerate(row):
            cell = gt.cell(ri, ci)
            cell.margin_left = Inches(0.06); cell.margin_right = Inches(0.06)
            cell.margin_top = Inches(0.02); cell.margin_bottom = Inches(0.02)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            p = cell.text_frame.paragraphs[0]
            _set(p, str(val), size,
                 WHITE if (header and ri == 0) else INK,
                 bold=(header and ri == 0),
                 align=PP_ALIGN.LEFT if ci == 0 else PP_ALIGN.CENTER)
            if header and ri == 0:
                cell.fill.solid(); cell.fill.fore_color.rgb = INK
            else:
                cell.fill.solid()
                cell.fill.fore_color.rgb = WHITE if ri % 2 else PANEL
    # strip default banding style
    tbl = gt._tbl
    for el in tbl.findall('.//{http://schemas.openxmlformats.org/drawingml/2006/main}tableStyleId'):
        el.text = '{5940675A-B579-460E-94D1-54222C63F5DA}'
    return gt


# ---- pipeline / flow boxes --------------------------------------------------
def flow_boxes(slide, x, y, w, h, steps, gap=Inches(0.12), vertical=False,
               fill=ACCENT2, size=10, head_size=11):
    """steps: list of (head, sub) tuples. Lay out as arrow-separated boxes."""
    n = len(steps)
    if vertical:
        bh = Emu(int((h - gap * (n - 1)) / n))
        for i, (head, sub) in enumerate(steps):
            by = Emu(int(y + i * (bh + gap)))
            _flow_one(slide, x, by, w, bh, head, sub, fill, size, head_size)
            if i < n - 1:
                _, tf = _tb(slide, x + w // 2 - Inches(0.15),
                            Emu(int(by + bh - Inches(0.02))), Inches(0.3), gap)
                _set(tf.paragraphs[0], "▼", 10, MUTE, align=PP_ALIGN.CENTER)
    else:
        bw = Emu(int((w - gap * (n - 1)) / n))
        for i, (head, sub) in enumerate(steps):
            bx = Emu(int(x + i * (bw + gap)))
            _flow_one(slide, bx, y, bw, h, head, sub, fill, size, head_size)
            if i < n - 1:
                _, tf = _tb(slide, Emu(int(bx + bw - Inches(0.03))),
                            y + h // 2 - Inches(0.12), gap + Inches(0.06), Inches(0.3))
                _set(tf.paragraphs[0], "▶", 9, MUTE, align=PP_ALIGN.CENTER)


def _flow_one(slide, x, y, w, h, head, sub, fill, size, head_size):
    box = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, h)
    box.adjustments[0] = 0.12
    box.fill.solid(); box.fill.fore_color.rgb = fill
    box.line.fill.background(); box.shadow.inherit = False
    tf = box.text_frame; tf.word_wrap = True
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    tf.margin_left = tf.margin_right = Inches(0.06)
    tf.margin_top = tf.margin_bottom = Inches(0.03)
    _set(tf.paragraphs[0], head, head_size, WHITE, bold=True, align=PP_ALIGN.CENTER,
         space_after=2)
    if sub:
        _set(tf.add_paragraph(), sub, size, RGBColor(0xE7, 0xEC, 0xF2),
             align=PP_ALIGN.CENTER)


def divider_slide(prs, kicker, title, blurb=None):
    s = _blank(prs)
    bg = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, EMU_W, EMU_H)
    bg.fill.solid(); bg.fill.fore_color.rgb = INK; bg.line.fill.background()
    bg.shadow.inherit = False
    bar = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.9), Inches(3.05),
                             Inches(2.6), Inches(0.06))
    bar.fill.solid(); bar.fill.fore_color.rgb = ACCENT; bar.line.fill.background()
    bar.shadow.inherit = False
    _, tf = _tb(s, Inches(0.9), Inches(2.0), Inches(11.5), Inches(1.0))
    _set(tf.paragraphs[0], kicker, 15, RGBColor(0xD8, 0x8A, 0x80), bold=True)
    _, tf = _tb(s, Inches(0.9), Inches(3.35), Inches(11.5), Inches(1.6))
    _set(tf.paragraphs[0], title, 32, WHITE, bold=True)
    if blurb:
        _, tf = _tb(s, Inches(0.9), Inches(4.7), Inches(11.2), Inches(1.5))
        _set(tf.paragraphs[0], blurb, 13, RGBColor(0xC9, 0xC9, 0xD6))
    return s
