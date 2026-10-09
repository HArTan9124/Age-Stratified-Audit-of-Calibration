"""Print bounding boxes + overlap / off-slide warnings for a pptx."""
import sys
from pptx import Presentation
from pptx.util import Emu

SW, SH = Emu(Inches := 12192000), Emu(6858000)  # 13.333 x 7.5 in

def inch(v): return round(v / 914400, 2)

prs = Presentation(sys.argv[1])
for si, slide in enumerate(prs.slides, 1):
    print(f"\n===== SLIDE {si} =====")
    boxes = []
    for sh in slide.shapes:
        if sh.left is None:
            continue
        l, t, w, h = sh.left, sh.top, sh.width or 0, sh.height or 0
        kind = sh.shape_type
        txt = ""
        if sh.has_text_frame:
            txt = " | ".join(p.text for p in sh.text_frame.paragraphs if p.text)[:60]
        r, b = l + w, t + h
        flags = []
        if l < -5000 or t < -5000 or r > SW + 5000 or b > SH + 5000:
            flags.append("OFF-SLIDE")
        print(f"  [{inch(l):>5}, {inch(t):>5}]  {inch(w):>5} x {inch(h):>5}  "
              f"r={inch(r):>5} b={inch(b):>5}  {str(kind):<22} {flags} {txt}")
        boxes.append((inch(l), inch(t), inch(r), inch(b), txt or str(kind)))
    # crude text/image overlap check (ignore full-bg rectangle)
    for i in range(len(boxes)):
        for j in range(i + 1, len(boxes)):
            a, c = boxes[i], boxes[j]
            if a[2] <= 13.34 and a[3] <= 7.51 and a[0] <= 0.01 and a[1] <= 0.01:
                continue
            ox = max(0, min(a[2], c[2]) - max(a[0], c[0]))
            oy = max(0, min(a[3], c[3]) - max(a[1], c[1]))
            if ox > 0.15 and oy > 0.15:
                area = ox * oy
                if area > 0.4:
                    print(f"    ~overlap {area:.1f} in^2 : '{a[4][:30]}'  <>  '{c[4][:30]}'")
